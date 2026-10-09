# hr/tests/test_staff_export.py
"""
The payroll Excel export must show every component assigned to staff,
not only the ones in the legacy fixed template.
"""
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase

from branches.models import Branch
from common.managers import set_current_tenant
from hr.models import BonusDeductionRequest, SalaryComponent, Staff, StaffPayInfo
from hr.services.staff_export import StaffPayrollExportService
from users.models import Tenant


class StaffPayrollExportTests(TestCase):

    def setUp(self):
        self.tenant = Tenant.objects.create(name="Export Tenant", slug="export-tenant")
        set_current_tenant(self.tenant)
        self.addCleanup(set_current_tenant, None)
        self.branch = Branch.objects.create(name="Main", code="EX", tenant=self.tenant)
        self.staff = Staff.objects.create(
            first_name="Ada", last_name="Obi", branch=self.branch, tenant=self.tenant,
            pension_provider="Stanbic IBTC", is_pension_exempt=True,
        )

    def _assign(self, name, component_type, amount):
        component = SalaryComponent.objects.create(
            name=name, component_type=component_type, default_amount=Decimal(amount),
            branch=self.branch, tenant=self.tenant,
        )
        StaffPayInfo.objects.create(
            staff=self.staff, component=component, amount=Decimal(amount),
            branch=self.branch, tenant=self.tenant,
        )

    def _export(self):
        buffer = StaffPayrollExportService(Staff.objects.filter(pk=self.staff.pk), period_label="MARCH 2026").generate()
        ws = openpyxl.load_workbook(buffer).active
        headers = [c.value for c in ws[3]]
        row = {h: c.value for h, c in zip(headers, ws[7])}
        return headers, row

    def test_custom_components_get_their_own_columns(self):
        self._assign("Basic Salary", "EARNING", "100000")
        self._assign("Responsibility Allowance", "EARNING", "20000")
        self._assign("Cooperative Contribution", "DEDUCTION", "5000")
        self._assign("Uniform Deduction", "DEDUCTION", "1500")

        headers, row = self._export()

        self.assertEqual(row["Responsibility Allowance"], 20000)
        self.assertEqual(row["Cooperative Contribution"], 5000)
        self.assertEqual(row["Uniform Deduction"], 1500)
        self.assertNotIn("Other Deductions", headers)
        # Custom earnings sit before Gross Salary and count towards it.
        self.assertLess(headers.index("Responsibility Allowance"), headers.index("Gross Salary"))
        self.assertEqual(row["Gross Salary"], 120000)
        # Custom deductions sit inside the deductions block and count in the total.
        self.assertLess(headers.index("Dev. Levy & Other"), headers.index("Cooperative Contribution"))
        self.assertLess(headers.index("Uniform Deduction"), headers.index("Total Deductions"))
        self.assertEqual(row["Net Pay"], row["Gross Salary"] - row["Total Deductions"])
        self.assertGreaterEqual(row["Total Deductions"], 6500)

    def test_template_layout_unchanged_without_custom_components(self):
        self._assign("Basic Salary", "EARNING", "100000")
        headers, row = self._export()
        self.assertEqual(headers[:9], [
            "Name", "Basic Salary", "Housing Allowance", "Transport Allowance",
            "Entertain.", "Utility", "Lunch", "Leav Allow.", "Gross Salary",
        ])
        self.assertEqual(row["PFA"], "Stanbic IBTC")

    def test_approved_one_time_requests_for_the_month_are_exported(self):
        self._assign("Basic Salary", "EARNING", "100000")
        user = get_user_model().objects.create_user(
            username="export-user", email="export@test.com", password="x",
            tenant=self.tenant, branch=self.branch,
        )
        fine = SalaryComponent.objects.create(
            name="Lateness Fine", component_type="DEDUCTION", default_amount=Decimal("0"),
            branch=self.branch, tenant=self.tenant,
        )

        def request(ref, month, status):
            BonusDeductionRequest.objects.create(
                reference_number=ref, staff=self.staff, component=fine,
                amount=Decimal("2500"), reason="Late", for_month=month, status=status,
                requested_by=user, owner=user, branch=self.branch, tenant=self.tenant,
            )

        request("BDR-EXP-1", date(2026, 3, 1), BonusDeductionRequest.APPROVED)
        request("BDR-EXP-2", date(2026, 3, 1), BonusDeductionRequest.PENDING)
        request("BDR-EXP-3", date(2026, 4, 1), BonusDeductionRequest.APPROVED)

        headers, row = self._export()

        # Only the approved March request counts, under its payslip label.
        self.assertEqual(row["Lateness Fine (One-time)"], 2500)
        self.assertNotIn("Lateness Fine", headers)
        self.assertLess(headers.index("Lateness Fine (One-time)"), headers.index("Total Deductions"))
        self.assertEqual(row["Net Pay"], row["Gross Salary"] - row["Total Deductions"])
        self.assertGreaterEqual(row["Total Deductions"], 2500)
